import os
import sys

def build_quotation_docx():
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn

    doc = docx.Document()

    # Page Margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Color Palette
    PRIMARY = RGBColor(30, 58, 138)      # Deep Navy #1E3A8A
    SECONDARY = RGBColor(14, 116, 144)   # Teal Cyan #0E7490
    ACCENT = RGBColor(217, 119, 6)       # Amber Gold #D97706
    DARK_TEXT = RGBColor(30, 41, 59)     # Charcoal #1E293B

    HEX_PRIMARY = "1E3A8A"
    HEX_LIGHT_BG = "F8FAFC"
    HEX_ACCENT_BG = "FEF3C7"
    HEX_BORDER = "CBD5E1"

    # Helper: Set Cell Shading
    def set_cell_bg(cell, hex_color):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    # Helper: Set Cell Margins (Padding)
    def set_cell_padding(cell, top=120, bottom=120, left=150, right=150):
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        cell._tc.get_or_add_tcPr().append(tcMar)

    # Helper: Add Custom Heading
    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = PRIMARY
        return p

    # 1. Header Table Banner
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_h = header_table.rows[0].cells[0]
    c_h.width = Inches(7.0)
    set_cell_bg(c_h, HEX_PRIMARY)
    set_cell_padding(c_h, top=200, bottom=200, left=200, right=200)

    p_main = c_h.paragraphs[0]
    p_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_main.paragraph_format.space_after = Pt(4)
    r_m = p_main.add_run("COMMANDER CITY SHUTTLE PASS SYSTEM")
    r_m.font.name = "Arial"
    r_m.font.size = Pt(18)
    r_m.font.bold = True
    r_m.font.color.rgb = RGBColor(255, 255, 255)

    p_sub = c_h.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(0)
    r_s = p_sub.add_run("SOFTWARE QUOTATION & MONTHLY MEMBERSHIP APPROVAL REQUEST")
    r_s.font.name = "Arial"
    r_s.font.size = Pt(11)
    r_s.font.bold = True
    r_s.font.color.rgb = RGBColor(245, 158, 11)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(6)

    # 2. Metadata Grid
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Document Ref:", "CC-SPS-QT-2026-09", "Date:", "September 02, 2026"),
        ("Project Name:", "Commander City Shuttle Pass System", "Capacity:", "Up to 10,000 Active Pass Holders"),
        ("Prepared For:", "Management & Board of Directors", "Prepared By:", "Software & IT Engineering Team"),
        ("Pricing Model:", "PKR 10,000 / Month Membership", "Status:", "Submitted for Management Approval")
    ]

    for row_idx, data in enumerate(meta_info):
        row = meta_table.rows[row_idx]
        bg = HEX_LIGHT_BG if row_idx % 2 == 0 else "FFFFFF"
        
        # Col 1
        c1 = row.cells[0]
        c1.width = Inches(3.5)
        set_cell_bg(c1, bg)
        set_cell_padding(c1, top=80, bottom=80, left=120, right=120)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1_lbl = p1.add_run(f"{data[0]} ")
        r1_lbl.font.bold = True
        r1_lbl.font.color.rgb = PRIMARY
        p1.add_run(data[1])

        # Col 2
        c2 = row.cells[1]
        c2.width = Inches(3.5)
        set_cell_bg(c2, bg)
        set_cell_padding(c2, top=80, bottom=80, left=120, right=120)
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2_lbl = p2.add_run(f"{data[2]} ")
        r2_lbl.font.bold = True
        r2_lbl.font.color.rgb = PRIMARY
        if row_idx == 3:
            r2_val = p2.add_run(data[3])
            r2_val.font.bold = True
            r2_val.font.color.rgb = ACCENT
        else:
            p2.add_run(data[3])

    # 3. Section 1: Executive Summary
    add_heading_1("1. Executive Summary")
    p_exec = doc.add_paragraph(
        "The Commander City Shuttle Pass System is a custom-built, enterprise transport management web software "
        "designed to digitize, automate, and centralize monthly shuttle pass renewals and passenger verification across Commander City. "
        "Engineered to manage up to 10,000 active pass holders, the software replaces manual paper registers with automated digital pass issuance, "
        "auto-formatted monthly serial codes, live Google Sheets synchronization, and 1-click batch renewals."
    )
    p_exec.paragraph_format.space_after = Pt(10)

    # 4. Section 2: Key Features & Plus Points
    add_heading_1("2. Core Plus-Point Features & Operational Value")

    features = [
        ("⚡ 1-Click Monthly Bulk Pass Renewal", 
         "Renews all active member passes for the next month in less than 3 seconds. Automatically updates start/end validity dates and increments auto-formatted serial numbers (e.g. SP-2026-09-001) without manual typing."),
        
        ("📊 Real-Time Google Sheets Auto-Sync", 
         "Features live cloud integration with Google Sheets. Administrators and finance teams can view real-time pass renewals, active counts, route usage, and payment records remotely from any smartphone or PC."),
        
        ("🖨️ Thermal & ID Card Batch Printing", 
         "Includes built-in thermal card and standard ID card layouts with passenger photos, CNIC numbers, route details, vehicle numbers, issue/expiry dates, and auto-generated verification serial barcodes."),
        
        ("💾 Offline & Online Hybrid Database (IndexedDB)", 
         "Operates 100% offline during internet outages using a browser-based IndexedDB engine. Data automatically queues locally and syncs to cloud servers once connectivity is restored, ensuring zero gate downtime."),
        
        ("👥 Complete Member Directory & Photo Repository", 
         "Centralized digital directory holding passenger profiles, contact numbers, emergency contacts, designated routes, and photo IDs for conductor verification."),
        
        ("📥 Excel & JSON Data Migration", 
         "Instant 1-click import functionality for existing Excel spreadsheets (.xlsx, .csv) and JSON database backups, allowing instant onboarding of historical records."),
        
        ("🛡️ Tamper-Proof Audit Control & Fraud Prevention", 
         "Standardized auto-serial generation prevents counterfeit, duplicate, or expired pass usage, eliminating revenue leakage for Commander City.")
    ]

    for title, desc in features:
        p_f = doc.add_paragraph()
        p_f.paragraph_format.space_before = Pt(2)
        p_f.paragraph_format.space_after = Pt(4)
        p_f.paragraph_format.left_indent = Inches(0.2)
        r_t = p_f.add_run(f"•  {title}: ")
        r_t.font.bold = True
        r_t.font.color.rgb = SECONDARY
        r_d = p_f.add_run(desc)
        r_d.font.color.rgb = DARK_TEXT

    # 5. Section 3: Commercial Proposal & Pricing
    add_heading_1("3. Commercial Proposal & Monthly Membership Pricing")
    p_pintro = doc.add_paragraph(
        "To maintain continuous system operations, Google Cloud synchronization, dedicated backups, "
        "and ongoing technical enhancements for up to 10,000 active members, we propose an all-inclusive monthly membership model:"
    )
    p_pintro.paragraph_format.space_after = Pt(8)

    # Pricing Table
    price_table = doc.add_table(rows=4, cols=3)
    price_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Service Module", "Scope & Capacity", "Monthly Fee (PKR)"]
    hdr_row = price_table.rows[0]
    for idx, text in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_bg(cell, HEX_PRIMARY)
        set_cell_padding(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if idx == 2:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(text)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    items = [
        ("Software Application License & Core System Access", "Up to 10,000 Active Pass Holders & Terminals", "INCLUDED"),
        ("Google Sheets Real-Time Cloud Auto-Sync & Storage", "Continuous 24/7 Live Sync & Backup Monitoring", "INCLUDED"),
        ("Monthly Maintenance & Technical Support Plan", "Dedicated On-Call Support & Routine Enhancements", "PKR 10,000 / Mo")
    ]

    for r_idx, (comp, scope, fee) in enumerate(items, start=1):
        row = price_table.rows[r_idx]
        bg_col = HEX_LIGHT_BG if r_idx % 2 == 1 else "FFFFFF"
        
        c1, c2, c3 = row.cells[0], row.cells[1], row.cells[2]
        c1.width = Inches(3.2)
        c2.width = Inches(2.3)
        c3.width = Inches(1.5)

        for c, val in zip([c1, c2, c3], [comp, scope, fee]):
            set_cell_bg(c, bg_col)
            set_cell_padding(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if c == c3:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = p.add_run(val)
                r.font.bold = True
                r.font.color.rgb = ACCENT if "10,000" in val else RGBColor(5, 150, 105)
            else:
                p.add_run(val)

    # Summary Total Callout Box
    p_tot = doc.add_paragraph()
    p_tot.paragraph_format.space_before = Pt(8)
    p_tot.paragraph_format.space_after = Pt(12)
    
    tot_table = doc.add_table(rows=1, cols=2)
    tot_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tr = tot_table.rows[0]
    
    c_tl = tr.cells[0]
    c_tl.width = Inches(4.5)
    set_cell_bg(c_tl, "E2E8F0")
    set_cell_padding(c_tl, top=100, bottom=100, left=120, right=120)
    p_tl = c_tl.paragraphs[0]
    p_tl.paragraph_format.space_after = Pt(0)
    r_tl = p_tl.add_run("TOTAL RECURRING MONTHLY MEMBERSHIP FEE:")
    r_tl.font.bold = True
    r_tl.font.color.rgb = PRIMARY

    c_tr = tr.cells[1]
    c_tr.width = Inches(2.5)
    set_cell_bg(c_tr, HEX_ACCENT_BG)
    set_cell_padding(c_tr, top=100, bottom=100, left=120, right=120)
    p_tr = c_tr.paragraphs[0]
    p_tr.paragraph_format.space_after = Pt(0)
    p_tr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_tr = p_tr.add_run("PKR 10,000 / Month")
    r_tr.font.bold = True
    r_tr.font.size = Pt(12)
    r_tr.font.color.rgb = ACCENT

    # 6. Section 4: Approval Request & Signatures
    add_heading_1("4. Formal Approval & Authorization Request")
    p_app = doc.add_paragraph(
        "Management approval is hereby requested to authorize system operations and sanction the recurring "
        "monthly membership fee of PKR 10,000 starting from September 2026."
    )
    p_app.paragraph_format.space_after = Pt(10)

    # Approval Sign-off Table
    app_table = doc.add_table(rows=5, cols=2)
    app_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    app_rows_data = [
        ("Approval Action Status:", "[  ] APPROVED       [  ] APPROVED WITH AMENDMENTS       [  ] REJECTED"),
        ("Approved Monthly Fee:", "PKR 10,000 / Month (Rupees Ten Thousand Only)"),
        ("Authorized Approver Name:", "____________________________________________________________________"),
        ("Designation / Title:", "____________________________________________________________________"),
        ("Approval Date & Stamp:", "________________________ , 2026")
    ]

    # Header Row
    c_ah1, c_ah2 = app_table.rows[0].cells[0], app_table.rows[0].cells[1]
    c_ah1.width, c_ah2.width = Inches(2.5), Inches(4.5)
    set_cell_bg(c_ah1, HEX_PRIMARY)
    set_cell_bg(c_ah2, HEX_PRIMARY)
    set_cell_padding(c_ah1, top=100, bottom=100, left=120, right=120)
    set_cell_padding(c_ah2, top=100, bottom=100, left=120, right=120)
    
    p_h1 = c_ah1.paragraphs[0]
    p_h1.paragraph_format.space_after = Pt(0)
    r_h1 = p_h1.add_run("Approval Parameter")
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(255, 255, 255)

    p_h2 = c_ah2.paragraphs[0]
    p_h2.paragraph_format.space_after = Pt(0)
    r_h2 = p_h2.add_run("Authorization Details")
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (label, val) in enumerate(app_rows_data, start=0):
        row = app_table.rows[idx]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width, c2.width = Inches(2.5), Inches(4.5)
        set_cell_bg(c1, HEX_LIGHT_BG)
        set_cell_bg(c2, "FFFFFF")
        set_cell_padding(c1, top=80, bottom=80, left=100, right=100)
        set_cell_padding(c2, top=80, bottom=80, left=100, right=100)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(label)
        r1.font.bold = True
        r1.font.color.rgb = PRIMARY

        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(val)
        if idx == 0:
            r2.font.bold = True
            r2.font.color.rgb = ACCENT

    p_space2 = doc.add_paragraph()
    p_space2.paragraph_format.space_after = Pt(16)

    # Signature Block
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    c_s1, c_s2 = sig_table.rows[0].cells[0], sig_table.rows[0].cells[1]
    c_s1.width, c_s2.width = Inches(3.5), Inches(3.5)

    p_s1 = c_s1.paragraphs[0]
    p_s1.paragraph_format.space_after = Pt(2)
    r_s1_head = p_s1.add_run("Submitted By:\n\n\n_____________________________________\n")
    r_s1_head.font.bold = True
    r_s1_head.font.color.rgb = PRIMARY
    r_s1_body = p_s1.add_run("Software Systems Lead / Admin\nCommander City Transport Dept.")
    r_s1_body.font.size = Pt(9.5)

    p_s2 = c_s2.paragraphs[0]
    p_s2.paragraph_format.space_after = Pt(2)
    r_s2_head = p_s2.add_run("Approved & Sanctioned By:\n\n\n_____________________________________\n")
    r_s2_head.font.bold = True
    r_s2_head.font.color.rgb = PRIMARY
    r_s2_body = p_s2.add_run("General Manager / Director Finance\nCommander City Executive Board")
    r_s2_body.font.size = Pt(9.5)

    # Save to target paths
    targets = [
        r"e:\SHUTTLE PASS PAYMENTS\QUOTATION.docx",
        r"e:\SHUTTLE PASS PAYMENTS\shuttle-pass-system\QUOTATION.docx",
        r"e:\SHUTTLE PASS PAYMENTS\shuttle-pass-system\QUOTATION_Commander_City_Shuttle_Pass_System.docx"
    ]

    for target in targets:
        os.makedirs(os.path.dirname(target), exist_ok=True)
        doc.save(target)
        print(f"Successfully generated DOCX with python-docx at: {target}")

if __name__ == "__main__":
    build_quotation_docx()
